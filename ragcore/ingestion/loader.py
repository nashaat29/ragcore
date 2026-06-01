from __future__ import annotations

from pathlib import Path

from ragcore.core.exceptions import IngestionError, UnsupportedFileTypeError
from ragcore.core.logger import get_logger

logger = get_logger("ragcore.ingestion.loader")

SUPPORTED_EXTENSIONS = {
    ".txt": "text",
    ".md": "markdown",
    ".pdf": "pdf",
    ".html": "html",
    ".htm": "html",
    ".docx": "docx",
    ".csv": "csv",
    ".json": "json",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".xml": "xml",
}


class Document:
    def __init__(self, page_content: str, metadata: dict | None = None):
        self.page_content = page_content
        self.metadata = metadata or {}

    def __repr__(self) -> str:
        return f"Document(metadata={self.metadata}, content_length={len(self.page_content)})"


class DocumentLoader:
    def __init__(self, file_path: str | Path):
        self.file_path = Path(file_path)
        self._validate()

    def _validate(self) -> None:
        if not self.file_path.exists():
            raise IngestionError(f"File not found: {self.file_path}")
        ext = self.file_path.suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise UnsupportedFileTypeError(
                f"Unsupported file type '{ext}'. Supported: {list(SUPPORTED_EXTENSIONS.keys())}"
            )

    def load(self) -> list[Document]:
        ext = self.file_path.suffix.lower()
        loader_name = SUPPORTED_EXTENSIONS[ext]
        loader_method = getattr(self, f"_load_{loader_name}", None)
        if loader_method is None:
            raise UnsupportedFileTypeError(f"No loader for type: {loader_name}")
        return loader_method()

    async def load_async(self) -> list[Document]:
        import asyncio

        return await asyncio.to_thread(self.load)

    def _load_text(self) -> list[Document]:
        content = self.file_path.read_text(encoding="utf-8", errors="replace")
        return [Document(page_content=content, metadata={"source": str(self.file_path), "type": "text"})]

    def _load_markdown(self) -> list[Document]:
        return self._load_text()

    def _load_html(self) -> list[Document]:
        try:
            from html.parser import HTMLParser

            class TextExtractor(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self._text: list[str] = []
                    self._skip = False

                def handle_starttag(self, tag, attrs):
                    if tag in ("script", "style", "nav", "footer", "header"):
                        self._skip = True

                def handle_endtag(self, tag):
                    if tag in ("script", "style", "nav", "footer", "header"):
                        self._skip = False

                def handle_data(self, data):
                    if not self._skip:
                        stripped = data.strip()
                        if stripped:
                            self._text.append(stripped)

                def get_text(self):
                    return "\n".join(self._text)

            parser = TextExtractor()
            content = self.file_path.read_text(encoding="utf-8", errors="replace")
            parser.feed(content)
            text = parser.get_text()
            return [Document(page_content=text, metadata={"source": str(self.file_path), "type": "html"})]
        except Exception as e:
            raise IngestionError(f"Failed to parse HTML: {e}") from e

    def _load_pdf(self) -> list[Document]:
        try:
            import pypdf

            docs: list[Document] = []
            reader = pypdf.PdfReader(str(self.file_path))
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                docs.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": str(self.file_path),
                            "type": "pdf",
                            "page": i + 1,
                            "total_pages": len(reader.pages),
                        },
                    )
                )
            return docs
        except ImportError:
            raise IngestionError("pypdf is required for PDF loading. Install with: pip install pypdf")
        except Exception as e:
            raise IngestionError(f"Failed to parse PDF: {e}") from e

    def _load_docx(self) -> list[Document]:
        try:
            from docx import Document as DocxDocument

            doc = DocxDocument(str(self.file_path))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            content = "\n".join(paragraphs)
            return [Document(page_content=content, metadata={"source": str(self.file_path), "type": "docx"})]
        except ImportError:
            raise IngestionError("python-docx is required for DOCX loading. Install with: pip install python-docx")
        except Exception as e:
            raise IngestionError(f"Failed to parse DOCX: {e}") from e

    def _load_csv(self) -> list[Document]:
        import csv
        import io

        content = self.file_path.read_text(encoding="utf-8", errors="replace")
        reader = csv.DictReader(io.StringIO(content))
        docs: list[Document] = []
        for i, row in enumerate(reader):
            text = "\n".join(f"{k}: {v}" for k, v in row.items() if v)
            docs.append(
                Document(
                    page_content=text,
                    metadata={"source": str(self.file_path), "type": "csv", "row": i + 1},
                )
            )
        return docs

    def _load_json(self) -> list[Document]:
        import json

        content = self.file_path.read_text(encoding="utf-8", errors="replace")
        data = json.loads(content)
        if isinstance(data, list):
            docs = []
            for i, item in enumerate(data):
                docs.append(
                    Document(
                        page_content=json.dumps(item, indent=2, default=str),
                        metadata={"source": str(self.file_path), "type": "json", "index": i},
                    )
                )
            return docs
        return [
            Document(
                page_content=json.dumps(data, indent=2, default=str),
                metadata={"source": str(self.file_path), "type": "json"},
            )
        ]

    def _load_yaml(self) -> list[Document]:
        import yaml

        content = self.file_path.read_text(encoding="utf-8", errors="replace")
        data = yaml.safe_load(content)
        text = yaml.dump(data, default_flow_style=False, allow_unicode=True) if data else content
        return [Document(page_content=text, metadata={"source": str(self.file_path), "type": "yaml"})]

    def _load_xml(self) -> list[Document]:
        import xml.etree.ElementTree as ET

        content = self.file_path.read_text(encoding="utf-8", errors="replace")
        root = ET.fromstring(content)

        def _iter_text(el: ET.Element) -> str:
            parts: list[str] = []
            if el.text and el.text.strip():
                parts.append(el.text.strip())
            for child in el:
                parts.append(_iter_text(child))
                if child.tail and child.tail.strip():
                    parts.append(child.tail.strip())
            return " ".join(parts)

        text = _iter_text(root)
        return [Document(page_content=text, metadata={"source": str(self.file_path), "type": "xml"})]


def load_directory(directory_path: str | Path, recursive: bool = True) -> list[Document]:
    base_path = Path(directory_path)
    if not base_path.is_dir():
        raise IngestionError(f"Not a directory: {directory_path}")

    pattern = "**/*" if recursive else "*"
    all_docs: list[Document] = []

    for file_path in sorted(base_path.glob(pattern)):
        if not file_path.is_file():
            continue
        ext = file_path.suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            continue
        try:
            loader = DocumentLoader(file_path)
            docs = loader.load()
            all_docs.extend(docs)
            logger.info(f"Loaded {file_path.name} -> {len(docs)} document(s)")
        except Exception as e:
            logger.warning(f"Skipped {file_path.name}: {e}")

    return all_docs


async def load_directory_async(directory_path: str | Path, recursive: bool = True) -> list[Document]:
    import asyncio

    return await asyncio.to_thread(load_directory, directory_path, recursive)
